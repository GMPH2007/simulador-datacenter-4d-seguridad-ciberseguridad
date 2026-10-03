import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_datacenter_case_study():
    doc = docx.Document()

    # Márgenes estándar (2.5 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Colores sobrios y formales
    NAVY = RGBColor(16, 44, 87)       # Azul corporativo títulos
    DARK_TEXT = RGBColor(33, 37, 41)   # Texto principal formal
    GRAY_TEXT = RGBColor(108, 117, 125)

    def set_cell_background(cell, color_hex):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_padding(cell, top=70, bottom=70, left=100, right=100):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for margin, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{margin}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def set_table_borders(table, border_hex="D1D5DB"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
            f'<w:left w:val="none"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    # ENCABEZADO FORMAL DEL TRABAJO
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = title_p.add_run("ACTIVIDAD DE APRENDIZAJE – SESIÓN DE CLASE")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(11)
    r_inst.font.bold = True
    r_inst.font.color.rgb = GRAY_TEXT

    main_title = doc.add_paragraph()
    main_title.paragraph_format.space_before = Pt(2)
    main_title.paragraph_format.space_after = Pt(12)
    main_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = main_title.add_run("SEGURIDAD FÍSICA Y LÓGICA EN UN DATA CENTER\nCASO PRÁCTICO Y MAPA CONCEPTUAL")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(13.5)
    r_title.font.bold = True
    r_title.font.color.rgb = NAVY

    # DATOS INFORMATIVOS (TABLA GRUPAL)
    info_table = doc.add_table(rows=4, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False
    set_table_borders(info_table, border_hex="E5E7EB")

    integrantes_lista = [
        "• Gerson Misael Pintado Huamán",
        "• Manuel Danilo López Garay",
        "• Mariana Juliet Clavijo Pinzón",
        "• Pedro Miguel Aguilar Flores",
        "• Dayron Antonio Urbina Zapata"
    ]

    info_data = [
        ("Integrantes:", integrantes_lista),
        ("Docente:", "Javier Eduardo Jaramillo Atoche"),
        ("Tema:", "1. Caso Práctico (Seguridad Física y Lógica) | 2. Mapa Conceptual"),
        ("Entorno / Caso:", "Centro de Cómputo y Data Center «FinanSur Data Cloud S.A.C.»")
    ]

    for i, (label, val) in enumerate(info_data):
        row = info_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(1.8)
        c1.width = Inches(4.7)
        set_cell_background(c0, "F3F4F6")
        set_cell_background(c1, "FFFFFF")
        set_cell_padding(c0, top=50, bottom=50, left=90, right=90)
        set_cell_padding(c1, top=50, bottom=50, left=90, right=90)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.name = "Arial"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = NAVY

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        if isinstance(val, list):
            for idx, item in enumerate(val):
                if idx > 0:
                    p1 = c1.add_paragraph()
                    p1.paragraph_format.space_before = Pt(1)
                    p1.paragraph_format.space_after = Pt(1)
                r1 = p1.add_run(item)
                r1.font.name = "Arial"
                r1.font.size = Pt(9.5)
                r1.font.color.rgb = DARK_TEXT
        else:
            r1 = p1.add_run(val)
            r1.font.name = "Arial"
            r1.font.size = Pt(9.5)
            r1.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # FUNCIONES HELPER
    def add_section_title(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = NAVY

    def add_subheading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.color.rgb = NAVY

    def add_p(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if bold_prefix:
            rb = p.add_run(bold_prefix + " ")
            rb.font.name = "Arial"
            rb.font.size = Pt(9.5)
            rb.font.bold = True
            rb.font.color.rgb = DARK_TEXT
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = DARK_TEXT
        return p

    def add_bullet(title, desc):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        rt = p.add_run(title + ": ")
        rt.font.name = "Arial"
        rt.font.size = Pt(9.5)
        rt.font.bold = True
        rt.font.color.rgb = NAVY
        rd = p.add_run(desc)
        rd.font.name = "Arial"
        rd.font.size = Pt(9.5)
        rd.font.color.rgb = DARK_TEXT

    # =========================================================================
    # PARTE 1: CASO PRÁCTICO
    # =========================================================================
    add_section_title("PARTE 1: CASO PRÁCTICO DE SEGURIDAD FÍSICA Y LÓGICA")
    add_subheading("1. Descripción de la Empresa y Situación Problemática")
    add_p(
        "La empresa «FinanSur Data Cloud S.A.C.» presta servicios de infraestructura tecnológica, procesamiento de pagos y alojamiento de bases de datos para entidades comerciales y bancarias de la región. Su centro de datos principal cuenta con 8 racks de servidores donde se almacenan datos financieros confidenciales de clientes."
    )
    add_p(
        "Durante el último mes, la empresa experimentó una contingencia grave: un corte del suministro eléctrico en la zona provocó una caída de servidores de 5 horas debido a que los sistemas UPS no tuvieron la autonomía suficiente y el generador secundario no arrancó automáticamente. Durante la revisión del incidente, se descubrió además que personal de limpieza externa había ingresado a la sala de servidores sin registro de bitácora, y que un intento de intrusión externa infectó uno de los servidores de desarrollo mediante un ataque de fuerza bruta por el puerto RDP expuesto a Internet sin cifrado ni doble factor de autenticación."
    )
    add_p(
        "Ante esta situación, la dirección general ordenó una auditoría inmediata para identificar todas las vulnerabilidades críticas y diseñar un plan integral de remediación que aborde tanto la Seguridad Física (instalaciones, energía y accesos) como la Seguridad Lógica (redes, sistemas y datos)."
    )

    # 2. DIAGNÓSTICO DE VULNERABILIDADES
    add_subheading("2. Diagnóstico de Vulnerabilidades Detectadas")
    add_p(
        "El informe de auditoría clasificó las fallas críticas encontradas en dos grupos técnicos bien definidos:"
    )

    diag_table = doc.add_table(rows=5, cols=2)
    diag_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    diag_table.autofit = False
    set_table_borders(diag_table)

    headers = ["Vulnerabilidades de Seguridad Física", "Vulnerabilidades de Seguridad Lógica"]
    for j, h_text in enumerate(headers):
        cell = diag_table.rows[0].cells[j]
        cell.width = Inches(3.25)
        set_cell_background(cell, "102C57")
        set_cell_padding(cell, top=70, bottom=70, left=90, right=90)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    brechas = [
        ("• Control de acceso deficiente: La puerta del data center utiliza cerradura convencional con llave física; no existe registro biométrico ni control de visitas.",
         "• Red plana sin segmentación (VLANs): Los servidores de base de datos se encuentran en la misma subred que las computadoras administrativas y WiFi de invitados."),
        ("• Climatización doméstica inadecuada: Se emplean aires acondicionados de tipo split convencionales que no controlan humedad, generando riesgo de condensación.",
         "• Ausencia de autenticación multifactor (MFA): Los administradores ingresan a los servidores principales únicamente con usuario y contraseña estática."),
        ("• Sistema contra incendios peligroso: Se mantienen rociadores de agua en el techo de la sala TI, lo que destruiría el hardware en caso de activación accidental.",
         "• Puertos administrativos expuestos: Puertos críticos (RDP 3389, SSH 22) abiertos directamente a Internet sin túneles VPN ni filtrado por listas blancas."),
        ("• Suministro eléctrico vulnerable: Banco de UPS monofásicos obsoletos sin redundancia N+1 y generador diésel con falla en el conmutador de transferencia (ATS).",
         "• Respaldos desprotegidos contra Ransomware: Las copias de seguridad se almacenan en un disco de red conectado permanentemente sin inmutabilidad ni copia fuera de línea.")
    ]

    for i, (col_f, col_l) in enumerate(brechas, start=1):
        bg = "F9FAFB" if i % 2 == 1 else "FFFFFF"
        row = diag_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = Inches(3.25), Inches(3.25)
        set_cell_background(c0, bg)
        set_cell_background(c1, bg)
        set_cell_padding(c0, top=60, bottom=60, left=80, right=80)
        set_cell_padding(c1, top=60, bottom=60, left=80, right=80)

        for c, txt in [(c0, col_f), (c1, col_l)]:
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r = p.add_run(txt)
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            r.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 3. PLAN DE SOLUCIONES TÉCNICAS
    add_subheading("3. Propuesta Integral de Soluciones Técnicas")
    add_p(
        "Para garantizar la disponibilidad, confidencialidad e integridad del Data Center, se estructura la solución en dos frentes simultáneos:"
    )

    add_p("A. Soluciones de Seguridad Física (Protección de Instalaciones e Infraestructura):", bold_prefix="[BLOQUE 1]")
    add_bullet(
        "Control de Acceso Biométrico y Esclusa de Paso (Mantrap)",
        "Instalación de un sistema de doble puerta (esclusa) que impide la apertura simultánea. El acceso requiere autenticación biométrica (huella/reconocimiento facial) y tarjeta de proximidad RFID cifrada, registrando automáticamente fecha, hora y usuario en una base de datos de auditoría."
    )
    add_bullet(
        "Climatización de Precisión y Confinamiento de Pasillos",
        "Reemplazo de los splits por dos unidades de aire acondicionado de precisión en configuración redundante (N+1). Se implementa confinamiento de pasillo frío para mantener la temperatura entre 18°C y 22°C y la humedad entre 40% y 60%, con sensores de alarma por fuga de agua."
    )
    add_bullet(
        "Sistema de Detección y Extinción por Gas Limpio (Novec 1230 / FM-200)",
        "Desconexión total de rociadores de agua en la sala técnica. Se instala un sistema de detección temprana por aspiración de humo (VESDA) y descarga automática de agente limpio que sofoca el fuego sin dejar residuos ni dañar los componentes electrónicos."
    )
    add_bullet(
        "Redundancia Eléctrica y Respaldo Continuo (Tier III)",
        "Instalación de un sistema de UPS trifásico modular de doble conversión con autonomía de 45 minutos y tablero de conmutación automática (ATS) conectado a un grupo electrógeno diésel con prueba de arranque semanal programada."
    )
    add_bullet(
        "Videovigilancia CCTV HD con Visión Nocturna",
        "Despliegue de cámaras IP con grabación continua 24/7 en pasillos de racks y accesos, con retención mínima de 90 días en almacenamiento protegido."
    )

    add_p("B. Soluciones de Seguridad Lógica (Protección de Redes, Datos y Sistemas):", bold_prefix="[BLOQUE 2]")
    add_bullet(
        "Segmentación de Red y Firewall Perimetral de Nueva Generación (NGFW)",
        "Configuración de VLANs dedicadas (VLAN Gestión, VLAN Servidores BD, VLAN Aplicaciones y DMZ). Se implementa un Firewall NGFW con inspección profunda de paquetes y políticas de denegación por defecto (Drop all by default)."
    )
    add_bullet(
        "Control de Acceso Lógico con Autenticación Multifactor (MFA/2FA)",
        "Exigencia obligatoria de token de seguridad de doble factor (MFA) para todo acceso administrativo a servidores (SSH, RDP, consolas web). Se aplica el principio de menor privilegio (RBAC) y desactivación inmediata de cuentas inactivas."
    )
    add_bullet(
        "Acceso Remoto Seguro mediante VPN Cifrada",
        "Cierre de todos los puertos directos a Internet. Cualquier tarea de mantenimiento remoto debe realizarse a través de un túnel VPN con cifrado IPsec o SSL/TLS robusto."
    )
    add_bullet(
        "Monitoreo Continuo con Sistema IDS/IPS y Registro de Logs (SIEM)",
        "Despliegue de sensores de Detección y Prevención de Intrusos (IDS/IPS) para bloquear escaneos de puertos y ataques web en tiempo real. Los logs de eventos se envían a un servidor SIEM centralizado para correlación de incidentes."
    )
    add_bullet(
        "Estrategia de Copias de Seguridad 3-2-1 con Inmutabilidad (Anti-Ransomware)",
        "Implementación de 3 copias de los datos críticos, en 2 medios distintos (disco local y almacenamiento cloud cifrado), con al menos 1 copia fuera de línea e inmutable (WORM - Write Once, Read Many), garantizando restauración íntegra ante desastres."
    )

    # 4. MATRIZ DE INTEGRACIÓN
    add_subheading("4. Matriz de Integración y Resultados Esperados")
    matrix_table = doc.add_table(rows=5, cols=4)
    matrix_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    matrix_table.autofit = False
    set_table_borders(matrix_table)

    m_headers = ["Amenaza o Fallo", "Medida Física Aplicada", "Medida Lógica Aplicada", "Impacto / Resultado"]
    col_w = [Inches(1.5), Inches(1.7), Inches(1.7), Inches(1.6)]

    for j, h_text in enumerate(m_headers):
        cell = matrix_table.rows[0].cells[j]
        cell.width = col_w[j]
        set_cell_background(cell, "102C57")
        set_cell_padding(cell, top=60, bottom=60, left=60, right=60)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    m_rows = [
        ("Acceso de personas no autorizadas al rack",
         "Esclusa de paso (Mantrap) + lector biométrico + CCTV 24/7.",
         "Cierre automático de sesiones y registro de accesos en SIEM.",
         "Cero accesos indebidos y trazabilidad del 100% de ingresos."),
        ("Caída del suministro eléctrico público",
         "UPS modular trifásico N+1 + Grupo electrógeno con ATS.",
         "Apagado ordenado automatizado por script si la batería agota.",
         "Disponibilidad 99.98% (SLA cumplido sin cortes imprevistos)."),
        ("Ataque de fuerza bruta y malware lateral",
         "Protección física de puertos de consola en racks con llave.",
         "Firewall NGFW, segmentación VLAN y MFA obligatorio en SSH/RDP.",
         "Bloqueo preventivo de intrusiones y aislamiento de amenazas."),
        ("Incendio o sobrecalentamiento en sala",
         "Aire de precisión N+1 y extinción por gas Novec 1230 sin agua.",
         "Sensores SNMP con alertas automáticas por correo y SMS.",
         "Continuidad operativa del hardware ante emergencias térmicas.")
    ]

    for i, row_data in enumerate(m_rows, start=1):
        bg = "F9FAFB" if i % 2 == 1 else "FFFFFF"
        row = matrix_table.rows[i]
        for j, text in enumerate(row_data):
            cell = row.cells[j]
            cell.width = col_w[j]
            set_cell_background(cell, bg)
            set_cell_padding(cell, top=50, bottom=50, left=60, right=60)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(8)
            r.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # =========================================================================
    # PARTE 2: MAPA CONCEPTUAL DEL TEMA
    # =========================================================================
    add_section_title("PARTE 2: MAPA CONCEPTUAL DEL TEMA")
    add_subheading("Esquema Gráfico Integral: Seguridad Física y Lógica en un Data Center")
    add_p(
        "A continuación se presenta el mapa conceptual que resume de forma jerárquica los componentes y controles indispensables para garantizar la seguridad en un centro de procesamiento de datos:"
    )

    # Insertar imagen generada si existe
    img_path = "mapa_conceptual_datacenter.png"
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(8)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(6.4))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(8)
        r_cap = p_cap.add_run("Figura 1: Mapa conceptual de Seguridad Física y Lógica en un Data Center")
        r_cap.font.name = "Arial"
        r_cap.font.size = Pt(8.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = GRAY_TEXT

    # Estructura textual complementaria del mapa conceptual
    add_subheading("Desglose Jerárquico del Mapa Conceptual")
    add_bullet(
        "1. Eje Central (Seguridad en el Data Center)",
        "Garantizar la tríada de seguridad de la información: Confidencialidad (solo usuarios autorizados), Integridad (datos exactos y no alterados) y Disponibilidad (sistemas operativos 24/7 sin interrupciones)."
    )
    add_bullet(
        "2. Dimensión de Seguridad Física (Hardware y Entorno)",
        "• Acceso: Esclusas (mantrap), biometría, tarjetas RFID y bitácora.\n• Vigilancia: Cámaras CCTV 24/7, sensores de apertura de racks e infrarrojos.\n• Clima: Aire acondicionado de precisión, pasillo frío/caliente, control de humedad.\n• Fuego: Detección VESDA y gas limpio Novec 1230 (sin agua).\n• Energía: Doble alimentación eléctrica, UPS modular y generador diésel con ATS."
    )
    add_bullet(
        "3. Dimensión de Seguridad Lógica (Datos, Redes y Software)",
        "• Redes: Firewalls NGFW, segmentación VLAN, DMZ y VPN cifrada.\n• Accesos: Autenticación Multifactor (MFA), política de menor privilegio (RBAC).\n• Monitoreo: Sistemas de prevención de intrusiones (IPS/IDS) y correlación SIEM.\n• Cifrado: En reposo (AES-256) y en tránsito (TLS 1.3).\n• Respaldos: Estrategia 3-2-1 con copias inmutables fuera de línea."
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # =========================================================================
    # PARTE 3: CONCLUSIONES
    # =========================================================================
    add_section_title("CONCLUSIONES GENERALES")
    add_bullet(
        "Interdependencia Indispensable",
        "La seguridad lógica carece de utilidad si la seguridad física falla; un atacante con acceso directo a un servidor puede reiniciar la máquina o extraer el disco en minutos. De igual modo, una instalación física blindada no resistirá si la red presenta puertos abiertos sin firewall ni MFA."
    )
    add_bullet(
        "Continuidad del Negocio y Resiliencia",
        "La infraestructura de un data center moderno debe concebirse bajo criterios de tolerancia a fallos (redundancia N+1 en energía y refrigeración) y copias de seguridad inmutables contra amenazas modernas como el Ransomware."
    )
    add_bullet(
        "Cultura de Cumplimiento y Auditoría",
        "Tanto las políticas físicas (registro de visitas, mantenimiento preventivo de UPS) como las lógicas (actualización de parches, revisión de logs SIEM) requieren auditorías periódicas para detectar brechas antes de que deriven en incidentes costosos."
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("________________________________________")
    r_end.font.color.rgb = GRAY_TEXT

    # Guardar en archivo
    output_filename = "Caso_Practico_Seguridad_DataCenter_Grupo.docx"
    doc.save(output_filename)
    print(f"Documento guardado exitosamente: {output_filename}")

if __name__ == "__main__":
    create_datacenter_case_study()
