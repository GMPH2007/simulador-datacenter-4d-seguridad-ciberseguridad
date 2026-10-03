import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_academic_case_study():
    doc = docx.Document()

    # Márgenes estándar para trabajos académicos (2.5 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Colores sobrios y formales
    NAVY = RGBColor(16, 44, 87)       # Azul formal para títulos
    DARK_TEXT = RGBColor(33, 37, 41)   # Texto principal formal
    GRAY_TEXT = RGBColor(108, 117, 125)

    def set_cell_background(cell, color_hex):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_padding(cell, top=80, bottom=80, left=120, right=120):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for margin, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{margin}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def set_table_borders(table, border_hex="D1D5DB"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
            f'<w:left w:val="none"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    # ENCABEZADO FORMAL DEL TRABAJO
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = title_p.add_run("ACTIVIDAD DE APRENDIZAJE N° 01")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(11)
    r_inst.font.bold = True
    r_inst.font.color.rgb = GRAY_TEXT

    main_title = doc.add_paragraph()
    main_title.paragraph_format.space_before = Pt(2)
    main_title.paragraph_format.space_after = Pt(12)
    main_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = main_title.add_run("CASO PRÁCTICO: SOLUCIONES EN HABILIDADES BLANDAS Y DURAS EN ENTORNOS DE EMPRESA")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(13.5)
    r_title.font.bold = True
    r_title.font.color.rgb = NAVY

    # DATOS INFORMATIVOS (TABLA GRUPAL)
    info_table = doc.add_table(rows=4, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False
    set_table_borders(info_table, border_hex="E5E7EB")

    integrantes_lista = [
        "• Gerson Misael Pintado Huamán",
        "• Manuel Danilo López Garay",
        "• Mariana Juliet Clavijo Pinzón",
        "• Pedro Miguel Aguilar Flores",
        "• Dayron Antonio Urbina Zapata"
    ]

    info_data = [
        ("Integrantes:", integrantes_lista),
        ("Docente:", "Javier Eduardo Jaramillo Atoche"),
        ("Tema:", "Propuesta de Soluciones en Habilidades Duras y Blandas"),
        ("Empresa Analizada:", "Distribuidora y Servicios Comerciales del Norte S.A.C.")
    ]

    for i, (label, val) in enumerate(info_data):
        row = info_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(1.8)
        c1.width = Inches(4.7)
        set_cell_background(c0, "F3F4F6")
        set_cell_background(c1, "FFFFFF")
        set_cell_padding(c0, top=60, bottom=60, left=100, right=100)
        set_cell_padding(c1, top=60, bottom=60, left=100, right=100)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.name = "Arial"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = NAVY

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        if isinstance(val, list):
            for idx, item in enumerate(val):
                if idx > 0:
                    p1 = c1.add_paragraph()
                    p1.paragraph_format.space_before = Pt(1)
                    p1.paragraph_format.space_after = Pt(1)
                r1 = p1.add_run(item)
                r1.font.name = "Arial"
                r1.font.size = Pt(9.5)
                r1.font.color.rgb = DARK_TEXT
        else:
            r1 = p1.add_run(val)
            r1.font.name = "Arial"
            r1.font.size = Pt(9.5)
            r1.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # FUNCIONES HELPER PARA TEXTO FORMAL
    def add_heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = NAVY

    def add_subheading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = NAVY

    def add_p(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if bold_prefix:
            rb = p.add_run(bold_prefix + " ")
            rb.font.name = "Arial"
            rb.font.size = Pt(9.5)
            rb.font.bold = True
            rb.font.color.rgb = DARK_TEXT
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = DARK_TEXT
        return p

    def add_bullet(title, desc):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        rt = p.add_run(title + ": ")
        rt.font.name = "Arial"
        rt.font.size = Pt(9.5)
        rt.font.bold = True
        rt.font.color.rgb = NAVY
        rd = p.add_run(desc)
        rd.font.name = "Arial"
        rd.font.size = Pt(9.5)
        rd.font.color.rgb = DARK_TEXT

    # 1. DESCRIPCIÓN DE LA EMPRESA Y SITUACIÓN PROBLEMÁTICA
    add_heading("1. Descripción de la Empresa y Situación Problemática")
    add_p(
        "La empresa «Distribuidora del Norte S.A.C.» se dedica a la comercialización y distribución de productos de consumo masivo para bodegas, minimarkets y restaurantes en la región. Su estructura operativa está conformada por tres áreas principales: Ventas (toma de pedidos), Almacén (preparación y despacho de mercadería) y Transporte (distribución final)."
    )
    add_p(
        "Durante los últimos cuatro meses, la gerencia general detectó un incremento considerable en los reclamos de clientes debido a pedidos incompletos, entregas fuera de horario y errores en las facturas y guías de remisión. Al mismo tiempo, el clima interno de la empresa se ha deteriorado: existen constantes desacuerdos entre el personal de almacén y los transportistas, acusaciones cruzadas de responsabilidades y un nivel de ausentismo laboral creciente provocado por el estrés en horas punta."
    )
    add_p(
        "La evaluación interna determinó que las causas no obedecen a falta de personal ni a fallas mecánicas de los vehículos, sino a dos deficiencias claras: debilidades técnicas en el uso de herramientas de control (habilidades duras) y serios problemas de comunicación, trabajo en equipo y liderazgo (habilidades blandas)."
    )

    # 2. IDENTIFICACIÓN Y DIAGNÓSTICO DE BRECHAS
    add_heading("2. Diagnóstico de Brechas en la Empresa")
    add_p(
        "A continuación, se clasifican las deficiencias detectadas en la operación diaria según correspondan al ámbito técnico o al ámbito interpersonal:"
    )

    diag_table = doc.add_table(rows=4, cols=2)
    diag_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    diag_table.autofit = False
    set_table_borders(diag_table)

    diag_headers = ["Deficiencias en Habilidades Duras (Técnicas)", "Deficiencias en Habilidades Blandas (Interpersonales)"]
    for j, h_text in enumerate(diag_headers):
        cell = diag_table.rows[0].cells[j]
        cell.width = Inches(3.25)
        set_cell_background(cell, "102C57")
        set_cell_padding(cell, top=70, bottom=70, left=90, right=90)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    diag_rows = [
        ("• Desconocimiento del sistema informático de inventarios: El personal registra movimientos de mercadería manualmente en hojas de cálculo desactualizadas, ocasionando diferencias entre el stock físico y el registrado.",
         "• Comunicación inasertiva: Cuando ocurre una equivocación en un pedido, los supervisores y trabajadores se comunican mediante reproches o gritos en lugar de orientar el diálogo hacia la solución."),
        ("• Falta de estandarización en el control de pedidos: No existen procedimientos operativos escritos para la verificación (picking y packing), lo que genera el envío de productos equivocados.",
         "• Falta de trabajo en equipo: Las áreas de Ventas, Almacén y Reparto operan como áreas aisladas. Ventas promete horarios de entrega imposibles sin consultar la capacidad operativa de almacén."),
        ("• Deficiente manejo de herramientas de cálculo (Excel / reportes): Los encargados no saben elaborar cuadros básicos para medir la cantidad de despachos diarios ni las mermas.",
         "• Nula tolerancia a la presión y mal manejo de conflictos: En los horarios de mayor salida de camiones, el personal se satura, pierde la calma y reacciona de forma agresiva ante los imprevistos.")
    ]

    for i, (col1, col2) in enumerate(diag_rows, start=1):
        bg = "F9FAFB" if i % 2 == 1 else "FFFFFF"
        row = diag_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = Inches(3.25), Inches(3.25)
        set_cell_background(c0, bg)
        set_cell_background(c1, bg)
        set_cell_padding(c0, top=60, bottom=60, left=80, right=80)
        set_cell_padding(c1, top=60, bottom=60, left=80, right=80)

        for c, txt in [(c0, col1), (c1, col2)]:
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r = p.add_run(txt)
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            r.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 3. PROPUESTA DE SOLUCIONES
    add_heading("3. Propuesta de Soluciones Integrales")
    add_p(
        "Para corregir los problemas identificados, se propone un plan de acción que combina capacitaciones técnicas específicas con talleres de desarrollo personal y clima laboral:"
    )

    add_subheading("A. Soluciones basadas en Habilidades Duras (Hard Skills)")
    add_bullet(
        "1. Capacitación práctica en el Sistema de Gestión de Inventarios",
        "Implementar un taller de 15 horas sobre el registro correcto de entradas, salidas y consulta de saldos en el software de la empresa. Se elaborarán manuales de una página plastificados y ubicados en los puestos de trabajo para consultas rápidas."
    )
    add_bullet(
        "2. Estandarización de procesos operativos (Picking y Packing)",
        "Crear una lista de chequeo (checklist) obligatoria antes de cargar cada camión. El operario y el chofer deben verificar juntos cantidad, código y estado del producto para evitar errores en destino."
    )
    add_bullet(
        "3. Capacitación en Excel Básico-Intermedio y Control de Métricas",
        "Enseñar a los jefes de grupo y supervisores a llevar el control diario de pedidos atendidos a tiempo, porcentaje de mercadería dañada y tiempo promedio de despacho por vehículo."
    )

    add_subheading("B. Soluciones basadas en Habilidades Blandas (Soft Skills)")
    add_bullet(
        "1. Taller de Comunicación Asertiva y Escucha Activa",
        "Capacitar al personal en formas correctas de expresar observaciones y pedidos sin caer en agresiones verbales ni actitudes pasivas. Se establecerá la regla de sustituir las quejas por propuestas de solución concretas."
    )
    add_bullet(
        "2. Coordinación y Trabajo en Equipo Interdepartamental",
        "Implementar una reunión breve de 10 minutos cada mañana entre un representante de Ventas, el supervisor de Almacén y el líder de Transporte para coordinar prioridades, revisar stock crítico y fijar compromisos claros del día."
    )
    add_bullet(
        "3. Taller de Manejo del Estrés y Resolución de Conflictos",
        "Brindar herramientas al personal para controlar la ansiedad en los momentos de alta carga laboral, enseñando pautas de negociación interna y respeto mutuo ante situaciones imprevistas."
    )
    add_bullet(
        "4. Desarrollo de Liderazgo Positivo para Supervisores",
        "Formar a los jefes de área en técnicas de retroalimentación constructiva, enseñándoles a corregir al colaborador en privado y con respeto, y a reconocer los logros del equipo de manera pública."
    )

    # 4. MATRIZ DE INTEGRACIÓN
    add_heading("4. Matriz de Integración: Relación entre Ambas Habilidades")
    add_p(
        "Las habilidades duras y blandas no actúan por separado; para resolver un problema real en la empresa se requiere aplicar ambas de forma simultánea, tal como se muestra en la siguiente tabla:"
    )

    matrix_table = doc.add_table(rows=4, cols=4)
    matrix_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    matrix_table.autofit = False
    set_table_borders(matrix_table)

    m_headers = ["Problema Operativo", "Habilidad Dura Requerida", "Habilidad Blanda Requerida", "Resultado Obtenido"]
    col_w = [Inches(1.5), Inches(1.7), Inches(1.7), Inches(1.6)]

    for j, h_text in enumerate(m_headers):
        cell = matrix_table.rows[0].cells[j]
        cell.width = col_w[j]
        set_cell_background(cell, "102C57")
        set_cell_padding(cell, top=60, bottom=60, left=60, right=60)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    m_rows = [
        ("Demora en el despacho de mercadería por desorden.",
         "Aplicación de orden técnico y codificación de pasillos (5S).",
         "Trabajo en equipo y cooperación entre almaceneros.",
         "Reducción del 40% en tiempos de preparación de carga."),
        ("Diferencias entre el inventario físico y el sistema.",
         "Conteo diario y conciliación técnica en planillas de control.",
         "Honestidad, ética y comunicación oportuna de errores.",
         "Exactitud de inventarios superior al 98%."),
        ("Reclamos de clientes por mercadería golpeada o errónea.",
         "Uso de listas de verificación (checklist) y normas de estiba.",
         "Empatía, actitud de servicio y buena atención.",
         "Disminución drástica de quejas y clientes fidelizados.")
    ]

    for i, row_data in enumerate(m_rows, start=1):
        bg = "F9FAFB" if i % 2 == 1 else "FFFFFF"
        row = matrix_table.rows[i]
        for j, text in enumerate(row_data):
            cell = row.cells[j]
            cell.width = col_w[j]
            set_cell_background(cell, bg)
            set_cell_padding(cell, top=50, bottom=50, left=60, right=60)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(8)
            r.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 5. CONCLUSIONES
    add_heading("5. Conclusiones")
    add_bullet(
        "Complementariedad en el desempeño laboral",
        "El conocimiento técnico por sí solo no garantiza el éxito de una empresa. Un trabajador puede saber operar el sistema a la perfección, pero si no sabe coordinar con sus compañeros o pierde la calma con los clientes, la operación se interrumpe."
    )
    add_bullet(
        "Importancia de la comunicación en la productividad",
        "La mayoría de los retrasos y errores en las empresas surgen por falta de diálogo claro entre departamentos. Fomentar la comunicación asertiva y el respeto mutuo previene fallas operativas y ahorra costos a la organización."
    )
    add_bullet(
        "Perfil del trabajador actual",
        "Las empresas valoran a los profesionales que demuestran tanto capacidad técnica (manejo de programas, herramientas y normas) como madurez emocional (puntualidad, compañerismo, adaptabilidad y actitud para resolver problemas)."
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("________________________________________")
    r_end.font.color.rgb = GRAY_TEXT

    # Guardamos en ambos nombres para que el alumno no tenga pérdida
    filenames = [
        "Caso_Practico_Habilidades_Blandas_y_Duras_Grupo.docx",
        "Caso_Practico_Habilidades_Blandas_y_Duras_Gerson_Pintado.docx"
    ]
    for fn in filenames:
        doc.save(fn)
        print(f"Documento guardado: {fn}")

if __name__ == "__main__":
    create_academic_case_study()
