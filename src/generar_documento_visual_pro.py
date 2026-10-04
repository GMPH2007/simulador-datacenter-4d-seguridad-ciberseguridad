import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def build_pro_max_document():
    doc = docx.Document()

    # Márgenes estándar equilibrados (2.2 cm)
    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    # Paleta de Colores Institucionales
    NAVY = RGBColor(10, 37, 64)         # #0A2540 - Azul profundo
    ROYAL_BLUE = RGBColor(29, 78, 216)  # #1D4ED8 - Azul acento
    TEAL = RGBColor(15, 118, 110)       # #0F766E - Verde azulado
    DARK_TEXT = RGBColor(30, 41, 59)    # #1E293B - Texto legible
    MUTED_TEXT = RGBColor(100, 116, 139) # #64748B - Subtítulos y pies

    # Helpers XML de formato
    def set_cell_background(cell, color_hex):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_padding(cell, top=70, bottom=70, left=100, right=100):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for margin, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{margin}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def set_table_borders(table, border_hex="CBD5E1"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
            f'<w:left w:val="none"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    # =========================================================================
    # CARÁTULA FORMAL PROFESIONAL (PÁGINA 1)
    # =========================================================================
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(25)
    p_inst.paragraph_format.space_after = Pt(2)
    r_inst = p_inst.add_run("INSTITUTO DE EDUCACIÓN SUPERIOR TECNOLÓGICO PÚBLICO")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(13)
    r_inst.font.bold = True
    r_inst.font.color.rgb = NAVY

    p_prog = doc.add_paragraph()
    p_prog.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_prog.paragraph_format.space_after = Pt(4)
    r_prog = p_prog.add_run("PROGRAMA DE ESTUDIOS: ARQUITECTURA DE PLATAFORMAS Y SERVICIOS DE TI")
    r_prog.font.name = "Arial"
    r_prog.font.size = Pt(11)
    r_prog.font.bold = True
    r_prog.font.color.rgb = ROYAL_BLUE

    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(35)
    r_div = p_div.add_run("____________________________________________________________")
    r_div.font.color.rgb = MUTED_TEXT

    p_tag = doc.add_paragraph()
    p_tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tag.paragraph_format.space_after = Pt(4)
    r_tag = p_tag.add_run("INFORME TÉCNICO DE APRENDIZAJE – SESIÓN DE CLASE")
    r_tag.font.name = "Arial"
    r_tag.font.size = Pt(11)
    r_tag.font.bold = True
    r_tag.font.color.rgb = MUTED_TEXT

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(8)
    r_title = p_title.add_run("SEGURIDAD FÍSICA Y LÓGICA EN UN DATA CENTER\nCASO PRÁCTICO INTEGRAL Y MAPA CONCEPTUAL INFOGRÁFICO")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(14.5)
    r_title.font.bold = True
    r_title.font.color.rgb = NAVY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(35)
    r_sub = p_sub.add_run("Aplicación de Estándares Internacionales: ANSI/TIA-942 (Tier III) e ISO/IEC 27001")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(10)
    r_sub.font.italic = True
    r_sub.font.color.rgb = ROYAL_BLUE

    # Ficha técnica de datos
    tbl_meta = doc.add_table(rows=3, cols=2)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_meta.autofit = False
    set_table_borders(tbl_meta, border_hex="94A3B8")

    integrantes_raw = (
        "• Gerson Misael Pintado Huamán\n"
        "• Manuel Danilo López Garay\n"
        "• Mariana Juliet Clavijo Pinzón\n"
        "• Pedro Miguel Aguilar Flores\n"
        "• Dayron Antonio Urbina Zapata"
    )

    meta_items = [
        ("DOCENTE:", "Dr. Javier Eduardo Jaramillo Atoche"),
        ("INTEGRANTES:", integrantes_raw),
        ("SEDE Y AÑO:", "Paita, Piura – Ciclo Lectivo 2026")
    ]

    for i, (label, val) in enumerate(meta_items):
        row = tbl_meta.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(1.9)
        c1.width = Inches(4.8)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF")
        set_cell_padding(c0, top=60, bottom=60, left=90, right=90)
        set_cell_padding(c1, top=60, bottom=60, left=90, right=90)

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.name = "Arial"
        r0.font.size = Pt(10)
        r0.font.bold = True
        r0.font.color.rgb = NAVY

        p1 = c1.paragraphs[0]
        for idx, line in enumerate(val.split('\n')):
            if idx > 0:
                p1 = c1.add_paragraph()
                p1.paragraph_format.space_before = Pt(1)
                p1.paragraph_format.space_after = Pt(1)
            r1 = p1.add_run(line)
            r1.font.name = "Arial"
            r1.font.size = Pt(10)
            r1.font.color.rgb = DARK_TEXT

    doc.add_page_break()

    # =========================================================================
    # HELPERS DE CONTENIDO
    # =========================================================================
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(12.5)
        r.font.bold = True
        r.font.color.rgb = NAVY

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = ROYAL_BLUE

    def add_p(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if bold_prefix:
            rb = p.add_run(bold_prefix + " ")
            rb.font.name = "Arial"
            rb.font.size = Pt(10)
            rb.font.bold = True
            rb.font.color.rgb = DARK_TEXT
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.color.rgb = DARK_TEXT
        return p

    def add_bullet(title, desc):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        rt = p.add_run(title + ": ")
        rt.font.name = "Arial"
        rt.font.size = Pt(10)
        rt.font.bold = True
        rt.font.color.rgb = NAVY
        rd = p.add_run(desc)
        rd.font.name = "Arial"
        rd.font.size = Pt(10)
        rd.font.color.rgb = DARK_TEXT

    def add_image_with_caption(img_path, caption_text, width_inches=6.4):
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(4)
            run = p_img.add_run()
            run.add_picture(img_path, width=Inches(width_inches))

            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(10)
            r_cap = p_cap.add_run(caption_text)
            r_cap.font.name = "Arial"
            r_cap.font.size = Pt(8.5)
            r_cap.font.italic = True
            r_cap.font.color.rgb = MUTED_TEXT

    # =========================================================================
    # ACTIVIDAD 1: CASO PRÁCTICO
    # =========================================================================
    add_h1("ACTIVIDAD 1: CASO PRÁCTICO DE SEGURIDAD FÍSICA Y LÓGICA EN UN DATA CENTER")
    
    add_h2("1.1 Presentación de la Empresa: «FinanSur Data Cloud S.A.C.»")
    add_p(
        "«FinanSur Data Cloud S.A.C.» es un centro de procesamiento de datos y hosting corporativo ubicado en la provincia de Paita. Administra una infraestructura crítica compuesta por 10 gabinetes (racks) de alta densidad, alojando bases de datos transaccionales, pasarelas de pago electrónico y sistemas de facturación de entidades financieras, cooperativas de ahorro y empresas agroexportadoras de la región."
    )
    add_p(
        "Por la naturaleza de sus operaciones, la empresa procesa diariamente más de 85,000 transacciones bancarias y comerciales. No obstante, una reciente auditoría y una serie de contingencias pusieron en evidencia que la infraestructura presentaba graves deficiencias de seguridad tanto física como lógica."
    )

    add_h2("1.2 Descripción del Incidente Real y Problemática Operativa")
    add_p(
        "El incidente detonante ocurrió el pasado 14 de agosto a las 11:20 a.m., originado por un apagón en el alimentador eléctrico industrial de la zona. La cadena de fallos se produjo en tres etapas críticas:"
    )
    add_bullet(
        "1. Falla Eléctrica y Colapso Térmico",
        "El banco de UPS monofásicos de la sala técnica, cuyas baterías no habían recibido mantenimiento preventivo en 18 meses, colapsó en apenas 7 minutos. Por su parte, el grupo electrógeno diésel exterior no arrancó de forma automática debido a que el conmutador de transferencia (ATS) había quedado en posición manual tras una revisión técnica previa. La sala quedó sin suministro y los aires acondicionados split domésticos se apagaron. En solo 25 minutos, la temperatura ambiente subió de 21°C a 39°C, activando el apagado de emergencia de los procesadores por sobrecalentamiento."
    )
    add_bullet(
        "2. Violación del Perímetro Físico y Accesos Indebidos",
        "Para acelerar la ventilación de la sala, los operadores abrieron de par en par la única puerta de acceso y colocaron ventiladores convencionales de pedestal. En medio de la confusión, personal de limpieza y proveedores externos de mensajería ingresaron a la sala técnica sin portar credencial, sin registro de bitácora y con libre acceso visual y físico frente a los racks descubiertos."
    )
    add_bullet(
        "3. Intrusión Lógica y Ejecución de Malware en Servidores",
        "Al reanudarse la energía, uno de los servidores de bases de datos reportó un consumo anormal del 100% de CPU y tráfico inusual hacia una IP externa. El análisis forense determinó que un ciberdelincuente aprovechó que el puerto de escritorio remoto RDP (3389) estaba expuesto a Internet sin VPN ni autenticación multifactor (MFA), ejecutando un ataque de diccionario sobre la cuenta genérica «Administrador» e inyectando un script malicioso en la memoria."
    )
    add_bullet(
        "4. Impacto Económico y de Reputación",
        "La caída del servicio se prolongó por 6 horas y media, provocando el bloqueo de transacciones comerciales, penalidades contractuales por S/. 48,000 y el reclamo formal de dos entidades bancarias socias."
    )

    # Insertamos la Imagen 1: Render de Seguridad Física
    add_image_with_caption(
        "datacenter_seguridad_fisica.jpg",
        "Figura 1: Diseño e infraestructura de Seguridad Física – Control de acceso en esclusa (Mantrap), biometría, sensores VESDA, Novec 1230 y confinamiento de pasillo frío."
    )

    add_h2("1.3 Cuestionario de Análisis del Caso Práctico")

    add_p("Pregunta 1: ¿Cuáles fueron las principales fallas de seguridad física identificadas?", bold_prefix="[Cuestionario]")
    add_p(
        "Respuesta: 1) Control de acceso precario (puerta con chapa simple sin esclusa de seguridad Mantrap ni biometría); 2) Inexistencia de bitácora digital de visitas; 3) Climatización doméstica inadecuada (equipos split sin control de humedad ni redundancia); 4) Suministro eléctrico frágil (UPS sin autonomía certificada, conmutador ATS trabado en manual); y 5) Rociadores de agua en sala TI con riesgo de destrucción de equipos."
    )

    add_p("Pregunta 2: ¿Cuáles fueron las principales brechas de seguridad lógica encontradas?", bold_prefix="[Cuestionario]")
    add_p(
        "Respuesta: 1) Red plana no segmentada (los servidores compartían el mismo dominio de difusión que las computadoras administrativas y el WiFi); 2) Exposición directa de puertos críticos (RDP 3389) hacia la Internet pública sin túnel VPN; 3) Uso de cuentas compartidas genéricas («Administrador»); 4) Carencia de autenticación de doble factor (MFA); 5) Inexistencia de copias de seguridad inmutables contra Ransomware; y 6) Falta de monitoreo SIEM para detección de anomalías."
    )

    add_p("Pregunta 3: ¿Qué marcos y normas internacionales deben normar la solución?", bold_prefix="[Cuestionario]")
    add_p(
        "Respuesta: Se fundamenta en dos normas globales: 1) ANSI/TIA-942 (Estándar de Infraestructura de Telecomunicaciones para Data Centers), adoptando el nivel Tier III con mantenimiento concurrente, rutas eléctricas redundantes y equipos N+1; y 2) ISO/IEC 27001 (Sistema de Gestión de Seguridad de la Información), implementando los controles del Anexo A para seguridad física del entorno (A.7) y seguridad operacional y de redes (A.8)."
    )

    # Insertamos la Imagen 2: Ciberseguridad Lógica / SOC
    add_image_with_caption(
        "datacenter_seguridad_logica.jpg",
        "Figura 2: Centro de Operaciones de Seguridad (SOC) – Monitoreo de Firewall NGFW, segmentación de VLANs, detección de intrusos IDS/IPS y cifrado activo AES-256."
    )

    add_h2("1.4 Plan Integral de Soluciones Técnicas Propuestas")

    add_p("A. Soluciones en Seguridad Física (Instalaciones y Hardware):", bold_prefix="[EJE FÍSICO]")
    add_bullet(
        "1. Esclusa de Seguridad Interbloqueada (Mantrap) y Biometría",
        "Instalación de un sistema de doble puerta electrónica donde la segunda puerta solo abre si la primera está completamente cerrada. El acceso exige tarjeta RFID cifrada y lector biométrico (huella dactilar y reconocimiento facial), con registro digital automático de fecha, hora y usuario."
    )
    add_bullet(
        "2. Climatización de Precisión con Confinamiento de Pasillos",
        "Reemplazo de los splits por unidades de aire de precisión en esquema redundante N+1. Se confina el pasillo frío para mantener la temperatura entre 18°C y 22°C y la humedad entre 40% y 60%, con sensores de alarma por inundación bajo el piso técnico."
    )
    add_bullet(
        "3. Detección Temprana (VESDA) y Gas Limpio (Novec 1230)",
        "Eliminación de rociadores de agua en la sala técnica. Se instala un sistema de detección temprana por aspiración de humo (VESDA) y extinción automática mediante gas limpio Novec 1230 o FM-200, que sofoca el fuego absorbiendo el calor sin dejar residuos ni conducir electricidad."
    )
    add_bullet(
        "4. Redundancia Eléctrica Certificada Tier III",
        "Doble acometida eléctrica (Rutas A y B), banco de UPS modular trifásico de doble conversión con 45 minutos de autonomía garantizada y conmutador de transferencia automática (ATS) conectado al generador diésel con prueba de encendido semanal programada."
    )
    add_bullet(
        "5. Videovigilancia CCTV HD con Visión Infrarroja 24/7",
        "Cámaras domo IP con sensor de movimiento y visión nocturna en pasillos y racks, con retención mínima de grabaciones por 90 días en almacenamiento protegido."
    )

    add_p("B. Soluciones en Seguridad Lógica (Redes, Sistemas y Datos):", bold_prefix="[EJE LÓGICO]")
    add_bullet(
        "1. Segmentación de Redes (VLANs / DMZ) y Firewall NGFW",
        "Despliegue de un Firewall de Nueva Generación en clúster activo-pasivo. Se segmenta la infraestructura en VLANs independientes (Gestión, Servidores BD, Aplicaciones DMZ y Usuarios), aplicando una política de denegación por defecto (Drop all by default)."
    )
    add_bullet(
        "2. Autenticación Multifactor (MFA/2FA) y Control RBAC",
        "Exigencia estricta de token de seguridad de segundo factor (MFA) para todo acceso a servidores e interfaces administrativas (SSH, RDP, consolas web). Se aplica el principio de menor privilegio (RBAC) con credenciales personales intransferibles."
    )
    add_bullet(
        "3. Acceso Remoto Seguro Exclusivo por VPN Cifrada",
        "Cierre total de los puertos de gestión hacia la red pública Internet. El soporte técnico a distancia se realiza únicamente a través de túneles VPN SSL/IPsec con cifrado AES-256."
    )
    add_bullet(
        "4. Detección de Intrusos (IDS/IPS) y Correlación SIEM",
        "Implementación de firmas de detección y prevención de intrusos (IDS/IPS) para bloquear escaneos de puertos y ataques web en tiempo real. Los registros de auditoría se centralizan en una plataforma SIEM con alertas 24/7."
    )
    add_bullet(
        "5. Estrategia de Copias de Seguridad 3-2-1 Inmutables",
        "3 copias de los datos críticos, en 2 soportes diferentes (almacenamiento local SAN y repositorio en la nube cifrado), con 1 copia fuera de línea e inmutable (WORM - Write Once, Read Many), garantizando la recuperación íntegra ante Ransomware."
    )

    add_h2("1.5 Matriz de Control de Riesgos y Resultados")
    
    tbl_mat = doc.add_table(rows=5, cols=4)
    tbl_mat.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_mat.autofit = False
    set_table_borders(tbl_mat)

    m_headers = ["Vulnerabilidad Detectada", "Control Físico Aplicado", "Control Lógico Aplicado", "Resultado e Impacto"]
    col_w = [Inches(1.5), Inches(1.7), Inches(1.7), Inches(1.6)]

    for j, h_text in enumerate(m_headers):
        cell = tbl_mat.rows[0].cells[j]
        cell.width = col_w[j]
        set_cell_background(cell, "0A2540")
        set_cell_padding(cell, top=60, bottom=60, left=60, right=60)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    m_rows = [
        ("Ingreso no autorizado de personas a la sala",
         "Esclusa Mantrap con biometría dactilar/facial y CCTV 24/7.",
         "Bloqueo automático de sesiones inactivas y log en SIEM.",
         "Cero ingresos indebidos; registro del 100% de accesos."),
        ("Corte eléctrico e interrupción del servicio",
         "UPS modular trifásico N+1 y generador diésel con ATS automático.",
         "Scripts de apagado ordenado automatizado si se agota la batería.",
         "Disponibilidad 99.98% (SLA cumplido sin caídas no planeadas)."),
        ("Ataque de fuerza bruta por puerto RDP",
         "Cerradura de seguridad en puertas frontales y posteriores de racks.",
         "Firewall NGFW, cierre de puertos, VPN obligatoria y MFA activo.",
         "Bloqueo total de conexiones directas; neutralización de intrusiones."),
        ("Riesgo de sobrecalentamiento e incendio",
         "Aire de precisión N+1 pasillo frío y extinción por gas Novec 1230.",
         "Monitoreo de temperatura por SNMP con alertas automáticas (SMS/Mail).",
         "Preservación íntegra del hardware ante emergencias térmicas.")
    ]

    for i, row_data in enumerate(m_rows, start=1):
        bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"
        row = tbl_mat.rows[i]
        for j, text in enumerate(row_data):
            cell = row.cells[j]
            cell.width = col_w[j]
            set_cell_background(cell, bg)
            set_cell_padding(cell, top=50, bottom=50, left=60, right=60)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(8)
            r.font.color.rgb = DARK_TEXT

    doc.add_page_break()

    # =========================================================================
    # ACTIVIDAD 2: MAPA CONCEPTUAL INFOGRÁFICO
    # =========================================================================
    add_h1("ACTIVIDAD 2: MAPA CONCEPTUAL DEL TEMA")
    
    add_h2("2.1 Representación Gráfica Infográfica del Mapa Conceptual")
    add_p(
        "A continuación se presenta el mapa conceptual en formato infográfico de alta definición, diseñado para ilustrar de forma clara, moderna y jerárquica la articulación de la Seguridad Física y de la Seguridad Lógica bajo el marco integral de protección de Centros de Datos:"
    )

    # Insertamos la Imagen 3: Infografía del Mapa Conceptual
    add_image_with_caption(
        "mapa_conceptual_infografia.jpg",
        "Figura 3: Mapa Conceptual Infográfico de Seguridad Física y Lógica en un Data Center (Marco Integral CID – ISO/IEC 27001 y ANSI/TIA-942 Tier III).",
        width_inches=6.6
    )

    add_h2("2.2 Estructura Jerárquica y Desglose por Niveles")
    add_p(
        "El mapa conceptual se articula a través de cuatro niveles jerárquicos que interconectan la teoría con la práctica operativa:"
    )

    map_levels = [
        ("Nivel 1 – Núcleo Conceptual:", "SEGURIDAD EN EL DATA CENTER. Representa el objetivo fundamental: salvaguardar los activos de hardware, software y comunicaciones garantizando la tríada CID (Confidencialidad, Integridad y Disponibilidad)."),
        ("Nivel 2 – Grandes Dimensiones:", "Se bifurca en dos ramas de igual criticidad: 1) SEGURIDAD FÍSICA (protección de activos tangibles, salas, accesos y factores ambientales) y 2) SEGURIDAD LÓGICA (protección de activos virtuales, redes, sistemas y datos contra ataques digitales)."),
        ("Nivel 3A – Pilares Físicos:", "Comprende cinco componentes de protección física: 1) Biometría y control perimetral; 2) Vigilancia CCTV 24/7; 3) Climatización HVAC de precisión; 4) Extinción por gas limpio; y 5) Respaldo eléctrico continuo."),
        ("Nivel 3B – Pilares Lógicos:", "Comprende cinco componentes de protección lógica: 1) Firewall NGFW perimetral; 2) Segmentación por VLANs; 3) Autenticación Multifactor (MFA); 4) Cifrado integral AES-256; y 5) Copias de seguridad en la nube (Cloud Backup)."),
        ("Nivel 4 – Mecanismos Tecnológicos:", "Implementa controles específicos como esclusas Mantrap, sensores VESDA, gas Novec 1230, UPS modular trifásico, ATS automático, políticas Drop All, tokens MFA, VPN IPsec y almacenamiento inmutable WORM.")
    ]

    for title, desc in map_levels:
        add_bullet(title, desc)

    add_h2("2.3 Glosario Técnico de Términos del Mapa Conceptual")
    
    tbl_glo = doc.add_table(rows=7, cols=2)
    tbl_glo.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_glo.autofit = False
    set_table_borders(tbl_glo)

    glosario_items = [
        ("Esclusa (Mantrap)", "Sistema de control de acceso físico con dos puertas interbloqueadas donde una no puede abrirse si la otra no está totalmente cerrada."),
        ("VESDA", "Sistema de detección de humo por aspiración de muy alta sensibilidad que detecta partículas microscópicas antes de que aparezca llama abierta."),
        ("Novec 1230 / FM-200", "Gases limpios de extinción que sofocan el conato de fuego absorbiendo el calor sin dejar residuos, sin ser conductores de electricidad y sin dañar la electrónica."),
        ("Conmutador ATS", "Interruptor de transferencia automática que detecta la caída de la energía pública y enciende el generador diésel transfiriendo la carga sin intervención humana."),
        ("Firewall NGFW", "Cortafuegos de nueva generación que analiza el tráfico a nivel de aplicación (Capa 7), integrando antivirus, prevención de intrusos (IPS) y filtrado web."),
        ("Autenticación MFA / 2FA", "Mecanismo de seguridad que exige al menos dos factores distintos para validar la identidad (algo que sabe: contraseña, y algo que tiene: token/código)."),
        ("Inmutabilidad WORM", "Propiedad de almacenamiento (Write Once, Read Many) que impide que una copia de respaldo sea modificada, sobreescrita o borrada por un ataque de Ransomware.")
    ]

    for i, (termino, def_txt) in enumerate(glosario_items):
        row = tbl_glo.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.0)
        c1.width = Inches(4.7)
        set_cell_background(c0, "F8FAFC")
        set_cell_background(c1, "FFFFFF")
        set_cell_padding(c0, top=50, bottom=50, left=80, right=80)
        set_cell_padding(c1, top=50, bottom=50, left=80, right=80)

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(termino)
        r0.font.name = "Arial"
        r0.font.size = Pt(9)
        r0.font.bold = True
        r0.font.color.rgb = NAVY

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(def_txt)
        r1.font.name = "Arial"
        r1.font.size = Pt(9)
        r1.font.color.rgb = DARK_TEXT

    # =========================================================================
    # CONCLUSIONES Y RECOMENDACIONES
    # =========================================================================
    add_h1("CONCLUSIONES Y RECOMENDACIONES GENERALES")
    
    conclusiones_finales = [
        ("Indisociabilidad de Ambas Dimensiones",
         "La seguridad física y la seguridad lógica son dos caras de la misma moneda. Blindar la red con firewalls y cifrado resulta estéril si cualquier persona puede ingresar a la sala y desconectar un servidor; de igual modo, un centro acorazado con biometría es inútil si los puertos de administración están abiertos a Internet sin protección."),
        ("Criterio de Redundancia y Continuidad (Tier III)",
         "En centros de cómputo que manejan operaciones críticas, no debe existir ningún punto único de falla (SPOF - Single Point of Failure). La duplicidad de rutas eléctricas (UPS y generadores) y la climatización N+1 son requisitos indispensables para cumplir con los acuerdos de nivel de servicio (SLA)."),
        ("Evolución hacia la Inmutabilidad de los Datos",
         "Frente al crecimiento de amenazas avanzadas como el Ransomware corporativo, las copias de seguridad tradicionales son insuficientes. Es mandatorio implementar esquemas de respaldo inmutables (WORM) y almacenamiento fuera de línea bajo la regla 3-2-1."),
        ("Mantenimiento y Auditoría Periódica",
         "La seguridad de un Data Center no concluye con la compra de equipos; exige pruebas periódicas de arranque de generadores, simulacros de contingencia, revisión semanal de bitácoras de acceso y auditorías continuas bajo la norma ISO/IEC 27001.")
    ]

    for title, desc in conclusiones_finales:
        add_bullet(title, desc)

    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_end.paragraph_format.space_before = Pt(25)
    r_end = p_end.add_run("— FIN DEL INFORME TÉCNICO —")
    r_end.font.name = "Arial"
    r_end.font.size = Pt(9.5)
    r_end.font.bold = True
    r_end.font.color.rgb = MUTED_TEXT

    # Guardamos en los archivos principales
    target_file = "TRABAJO_DATA_CENTER_PRO_MAX.docx"
    doc.save(target_file)
    print(f"Documento guardado: {target_file}")

    try:
        doc.save("Caso_Practico_y_Mapa_Conceptual_DataCenter_Grupo.docx")
        print("Documento guardado: Caso_Practico_y_Mapa_Conceptual_DataCenter_Grupo.docx")
    except Exception as e:
        print(f"Nota de bloqueo: {e}")

if __name__ == "__main__":
    build_pro_max_document()
