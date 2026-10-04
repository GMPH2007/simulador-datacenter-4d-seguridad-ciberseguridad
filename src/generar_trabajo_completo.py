import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def build_full_assignment():
    doc = docx.Document()

    # Configuración de márgenes estándar
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Colores sobrios institucionales
    NAVY = RGBColor(10, 37, 64)       # Azul noche formal
    BLUE_ACCENT = RGBColor(30, 64, 175) # Azul medio
    DARK_TEXT = RGBColor(33, 37, 41)   # Texto principal
    GRAY_TEXT = RGBColor(100, 116, 139) # Texto secundario / notas
    LIGHT_BG = "F8FAFC"
    HEADER_BG = "0A2540"

    def set_cell_background(cell, color_hex):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_padding(cell, top=70, bottom=70, left=100, right=100):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for margin, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{margin}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def set_table_borders(table, border_hex="D1D5DB"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
            f'<w:left w:val="none"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    # =========================================================================
    # PÁGINA 1: CARÁTULA OFICIAL FORMAL
    # =========================================================================
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(30)
    p_inst.paragraph_format.space_after = Pt(2)
    r_inst = p_inst.add_run("INSTITUTO DE EDUCACIÓN SUPERIOR TECNOLÓGICO PÚBLICO")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(13)
    r_inst.font.bold = True
    r_inst.font.color.rgb = NAVY

    p_carr = doc.add_paragraph()
    p_carr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_carr.paragraph_format.space_before = Pt(2)
    p_carr.paragraph_format.space_after = Pt(4)
    r_carr = p_carr.add_run("PROGRAMA DE ESTUDIOS: ARQUITECTURA DE PLATAFORMAS Y SERVICIOS DE TI")
    r_carr.font.name = "Arial"
    r_carr.font.size = Pt(11)
    r_carr.font.bold = True
    r_carr.font.color.rgb = BLUE_ACCENT

    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_after = Pt(40)
    r_line = p_line.add_run("____________________________________________________________")
    r_line.font.color.rgb = GRAY_TEXT

    # Título central de la actividad
    p_act = doc.add_paragraph()
    p_act.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_act.paragraph_format.space_after = Pt(4)
    r_act = p_act.add_run("ACTIVIDAD DE APRENDIZAJE – DESARROLLO DE SESIÓN")
    r_act.font.name = "Arial"
    r_act.font.size = Pt(12)
    r_act.font.bold = True
    r_act.font.color.rgb = GRAY_TEXT

    p_tit = doc.add_paragraph()
    p_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tit.paragraph_format.space_before = Pt(4)
    p_tit.paragraph_format.space_after = Pt(6)
    r_tit = p_tit.add_run("1. CASO PRÁCTICO: SEGURIDAD FÍSICA Y LÓGICA EN UN DATA CENTER\n2. MAPA CONCEPTUAL DEL TEMA")
    r_tit.font.name = "Arial"
    r_tit.font.size = Pt(15)
    r_tit.font.bold = True
    r_tit.font.color.rgb = NAVY

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(40)

    # Cuadro de datos
    tbl_caratula = doc.add_table(rows=3, cols=2)
    tbl_caratula.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_caratula.autofit = False
    set_table_borders(tbl_caratula, border_hex="CBD5E1")

    integrantes_txt = (
        "• Gerson Misael Pintado Huamán\n"
        "• Manuel Danilo López Garay\n"
        "• Mariana Juliet Clavijo Pinzón\n"
        "• Pedro Miguel Aguilar Flores\n"
        "• Dayron Antonio Urbina Zapata"
    )

    caratula_data = [
        ("DOCENTE:", "Dr. Javier Eduardo Jaramillo Atoche"),
        ("INTEGRANTES:", integrantes_txt),
        ("SEDE Y FECHA:", "Paita, Piura – Ciclo Lectivo 2026")
    ]

    for i, (etiqueta, valor) in enumerate(caratula_data):
        row = tbl_caratula.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.0)
        c1.width = Inches(4.7)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF")
        set_cell_padding(c0, top=60, bottom=60, left=90, right=90)
        set_cell_padding(c1, top=60, bottom=60, left=90, right=90)

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(etiqueta)
        r0.font.name = "Arial"
        r0.font.size = Pt(10)
        r0.font.bold = True
        r0.font.color.rgb = NAVY

        p1 = c1.paragraphs[0]
        for line_idx, line in enumerate(valor.split('\n')):
            if line_idx > 0:
                p1 = c1.add_paragraph()
                p1.paragraph_format.space_before = Pt(1)
                p1.paragraph_format.space_after = Pt(1)
            r1 = p1.add_run(line)
            r1.font.name = "Arial"
            r1.font.size = Pt(10)
            r1.font.color.rgb = DARK_TEXT

    # Salto de página para el contenido
    doc.add_page_break()

    # =========================================================================
    # HELPERS DE CONTENIDO
    # =========================================================================
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(12.5)
        r.font.bold = True
        r.font.color.rgb = NAVY

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = BLUE_ACCENT

    def add_p(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if bold_prefix:
            rb = p.add_run(bold_prefix + " ")
            rb.font.name = "Arial"
            rb.font.size = Pt(10)
            rb.font.bold = True
            rb.font.color.rgb = DARK_TEXT
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.color.rgb = DARK_TEXT
        return p

    def add_bullet(title, desc):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        rt = p.add_run(title + ": ")
        rt.font.name = "Arial"
        rt.font.size = Pt(10)
        rt.font.bold = True
        rt.font.color.rgb = NAVY
        rd = p.add_run(desc)
        rd.font.name = "Arial"
        rd.font.size = Pt(10)
        rd.font.color.rgb = DARK_TEXT

    # =========================================================================
    # ACTIVIDAD 1: CASO PRÁCTICO
    # =========================================================================
    add_h1("ACTIVIDAD 1: CASO PRÁCTICO DE SEGURIDAD FÍSICA Y LÓGICA EN UN DATA CENTER")
    
    add_h2("1.1 Presentación del Caso: Data Center de «FinanSur Data Cloud S.A.C.»")
    add_p(
        "«FinanSur Data Cloud S.A.C.» es una compañía proveedora de infraestructura informática y pasarela de pagos electrónicos para entidades bancarias, cooperativas y empresas de exportación de la zona norte. Su centro de datos principal administra 10 racks con servidores de aplicaciones, bases de datos transaccionales y almacenamiento en red (SAN/NAS)."
    )
    add_p(
        "El Data Center procesa un promedio de 85,000 transacciones diarias. Hasta hace tres meses, la empresa operaba bajo un esquema informal de seguridad que dependía de la confianza en el personal y de configuraciones básicas de fábrica, sin apego a estándares formales como ISO/IEC 27001 ni ANSI/TIA-942."
    )

    add_h2("1.2 Descripción del Incidente y Problemática Operativa")
    add_p(
        "El pasado 14 de agosto a las 11:20 a.m. se produjo un corte intempestivo del suministro eléctrico en la zona industrial. De inmediato se desencadenaron las siguientes contingencias críticas:"
    )
    add_bullet(
        "Fallo de Energía y Contingencia Térmica",
        "El banco de UPS monofásicos, con baterías degradadas por falta de mantenimiento, solo soportó 7 minutos de carga antes de apagarse. El grupo electrógeno diésel exterior no arrancó de forma automática porque el conmutador de transferencia (ATS) se encontraba en modo manual desde una prueba previa no reportada. La sala quedó a oscuras y los aires acondicionados convencionales de pared se detuvieron, provocando una subida de temperatura de 21°C a 39°C en 25 minutos, activando las alarmas térmicas de los procesadores."
    )
    add_bullet(
        "Violación de Seguridad Física Perimetral",
        "Durante la desesperación del corte, las puertas de la sala técnica se mantuvieron abiertas para ventilar el calor con ventiladores de pedestal. Personal de limpieza y contratistas externos ingresaron a la sala de servidores sin registro de bitácora ni gafete identificatorio, transitando libremente frente a los racks."
    )
    add_bullet(
        "Incidente de Seguridad Lógica y Acceso No Autorizado",
        "Al restablecerse el fluido eléctrico, se detectó que un servidor de base de datos presentaba un consumo inusitado del 100% de CPU. El análisis forense reveló que un atacante externo había aprovechado que el puerto RDP (3389) estaba abierto directamente hacia Internet sin VPN ni autenticación multifactor (MFA). El atacante ejecutó un ataque de fuerza bruta sobre la cuenta genérica «Administrador», logró acceso e inyectó un script malicioso en la memoria antes de ser desconectado."
    )
    add_bullet(
        "Consecuencias Económicas y Operativas",
        "La caída total de los servicios se extendió por 6 horas y media, provocando el retraso de liquidaciones bancarias, penalidades contractuales por S/. 48,000 y pérdida de confianza por parte de dos clientes corporativos."
    )

    add_h2("1.3 Cuestionario de Análisis del Caso Práctico")
    add_p("A continuación, se resuelven las preguntas de diagnóstico y evaluación técnica del caso:")

    add_p(
        "¿Cuáles fueron las principales fallas de seguridad física identificadas?",
        bold_prefix="Pregunta 1:"
    )
    add_p(
        "Respuesta: 1) Falta de control de acceso estricto (puertas abiertas sin esclusa de seguridad ni biometría); 2) Ausencia de bitácora obligatoria de visitas; 3) Climatización doméstica inadecuada (equipos split sin control de humedad ni redundancia); 4) Sistema de energía vulnerable (UPS sin autonomía certificada, generador sin conmutación automática ATS funcional); y 5) Rociadores de agua instalados sobre racks que ponían en riesgo de cortocircuito los equipos."
    )

    add_p(
        "¿Cuáles fueron las principales brechas de seguridad lógica encontradas?",
        bold_prefix="Pregunta 2:"
    )
    add_p(
        "Respuesta: 1) Ausencia de segmentación de red (los servidores compartían la misma subred que las PCs de oficina y el WiFi); 2) Exposición directa de puertos críticos de gestión (RDP 3389) a la red pública Internet sin túnel VPN; 3) Empleo de cuentas genéricas («Administrador») con contraseñas vulnerables; 4) Inexistencia de autenticación multifactor (MFA/2FA); 5) Inexistencia de copias de seguridad inmutables protegidas contra malware; y 6) Falta de un sistema centralizado de logs (SIEM) para detectar intrusiones en tiempo real."
    )

    add_p(
        "¿Qué normas y estándares internacionales deben aplicarse para solucionar el problema?",
        bold_prefix="Pregunta 3:"
    )
    add_p(
        "Respuesta: Se deben adoptar dos marcos fundamentales: 1) ANSI/TIA-942 (Estándar de Infraestructura de Telecomunicaciones para Centros de Datos), adoptando un nivel de redundancia Tier III (mantenimiento concurrente, componentes N+1 y rutas de energía redundantes); y 2) ISO/IEC 27001 (Sistema de Gestión de Seguridad de la Información), implementando los controles del Anexo A sobre seguridad física del entorno (A.7) y seguridad de las operaciones y comunicaciones (A.8)."
    )

    add_h2("1.4 Plan Integral de Soluciones Técnicas Propuestas")
    add_p("El equipo de trabajo plantea la siguiente remediación estructurada en dos dimensiones:")

    add_p("A. Soluciones en Seguridad Física:", bold_prefix="[DIMENSIÓN 1]")
    add_bullet(
        "Control de Acceso Biométrico y Esclusa Mantrap",
        "Implementación de un sistema de doble puerta interbloqueada (Mantrap) que impide que ambas puertas se abran a la vez. El ingreso exige identificación por tarjeta RFID cifrada y biometría dactilar o facial, registrando automáticamente el historial de accesos en una base de datos inalterable."
    )
    add_bullet(
        "Climatización de Precisión y Confinamiento de Pasillos",
        "Instalación de dos unidades de aire acondicionado de precisión con configuración redundante N+1. Se implementa confinamiento de pasillo frío para garantizar temperatura entre 18°C y 22°C y humedad entre 40% y 60%, con sensores de alarma por inundación bajo el piso técnico."
    )
    add_bullet(
        "Protección Contra Incendios con Agente Limpio (Novec 1230)",
        "Desmantelamiento de los rociadores de agua en la sala técnica. Se instala un sistema de detección temprana por aspiración de humo (VESDA) con extinción automática mediante gas limpio Novec 1230 o FM-200, inocuo para los equipos electrónicos y sin residuo."
    )
    add_bullet(
        "Infraestructura Eléctrica Redundante (Tier III)",
        "Incorporación de dos acometidas eléctricas (A y B), banco de UPS trifásico modular de doble conversión con autonomía de 45 minutos y conmutador ATS certificado que arranca el grupo electrógeno en menos de 10 segundos ante fallas de la red pública."
    )
    add_bullet(
        "Videovigilancia CCTV HD 24/7",
        "Instalación de cámaras IP fijas y domos en todos los ángulos de la sala y accesos, con grabación continua y retención mínima de 90 días en servidor de video aislado."
    )

    add_p("B. Soluciones en Seguridad Lógica:", bold_prefix="[DIMENSIÓN 2]")
    add_bullet(
        "Segmentación de Redes (VLANs / DMZ) y Firewall NGFW",
        "Despliegue de un Firewall de Nueva Generación en clúster activo-pasivo. Se segmenta la red mediante VLANs dedicadas: VLAN Gestión, VLAN Base de Datos, VLAN Servidores Web (DMZ) y VLAN Usuarios, bloqueando todo tráfico entre segmentos por defecto."
    )
    add_bullet(
        "Gestión de Identidades (IAM) y Autenticación Multifactor (MFA)",
        "Obligatoriedad de token de segundo factor (MFA) para todo acceso a consolas de administración. Se aplica el principio de menor privilegio (RBAC), eliminando las cuentas compartidas y exigiendo credenciales nominativas."
    )
    add_bullet(
        "Acceso Remoto Exclusivo por VPN Cifrada",
        "Cierre de todos los puertos directos a Internet (RDP, SSH, SMB). Todo acceso para soporte técnico remoto se canaliza a través de una VPN SSL con túnel cifrado AES-256."
    )
    add_bullet(
        "Detección de Intrusos (IDS/IPS) y Correlación de Eventos (SIEM)",
        "Activación de firmas de prevención de intrusos para bloquear escaneos de puertos y ataques web automáticamente. Los registros de todos los servidores y firewalls se envían a un servidor SIEM para monitoreo 24/7."
    )
    add_bullet(
        "Estrategia de Backup 3-2-1 con Inmutabilidad (Anti-Ransomware)",
        "Se generan 3 copias de los datos críticos, en 2 medios distintos (almacenamiento local SAN y repositorio en la nube cifrado), con 1 copia fuera de línea e inmutable (WORM), garantizando la recuperación ante incidentes de secuestro de datos."
    )

    add_h2("1.5 Matriz de Control de Riesgos y Resultados Esperados")
    
    matrix_tbl = doc.add_table(rows=5, cols=4)
    matrix_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    matrix_tbl.autofit = False
    set_table_borders(matrix_tbl)

    m_headers = ["Vulnerabilidad Detectada", "Control Físico Aplicado", "Control Lógico Aplicado", "Resultado e Impacto"]
    col_w = [Inches(1.5), Inches(1.7), Inches(1.7), Inches(1.6)]

    for j, h_text in enumerate(m_headers):
        cell = matrix_tbl.rows[0].cells[j]
        cell.width = col_w[j]
        set_cell_background(cell, HEADER_BG)
        set_cell_padding(cell, top=60, bottom=60, left=60, right=60)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    m_rows = [
        ("Ingreso no autorizado de personas",
         "Esclusa Mantrap con biometría dactilar/facial y CCTV 24/7.",
         "Bloqueo automático de sesiones inactivas y log en SIEM.",
         "Cero ingresos indebidos; registro del 100% de accesos."),
        ("Corte de energía e interrupción de servicios",
         "UPS modular trifásico N+1 y generador diésel con ATS automático.",
         "Scripts de apagado ordenado automatizado si se agota la batería.",
         "Disponibilidad 99.98% (SLA cumplido sin caídas no planeadas)."),
        ("Ataque de fuerza bruta por puerto RDP",
         "Cerradura de seguridad en puertas frontales y posteriores de racks.",
         "Firewall NGFW, cierre de puertos, VPN obligatoria y MFA activo.",
         "Bloqueo total de conexiones directas; neutralización de intrusiones."),
        ("Riesgo de sobrecalentamiento e incendio",
         "Aire de precisión N+1 pasillo frío y extinción por gas Novec 1230.",
         "Monitoreo de temperatura por SNMP con alertas automáticas (SMS/Mail).",
         "Preservación íntegra del hardware ante emergencias térmicas.")
    ]

    for i, row_data in enumerate(m_rows, start=1):
        bg = LIGHT_BG if i % 2 == 1 else "FFFFFF"
        row = matrix_tbl.rows[i]
        for j, text in enumerate(row_data):
            cell = row.cells[j]
            cell.width = col_w[j]
            set_cell_background(cell, bg)
            set_cell_padding(cell, top=50, bottom=50, left=60, right=60)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(8)
            r.font.color.rgb = DARK_TEXT

    doc.add_page_break()

    # =========================================================================
    # ACTIVIDAD 2: MAPA CONCEPTUAL DEL TEMA
    # =========================================================================
    add_h1("ACTIVIDAD 2: MAPA CONCEPTUAL DE SEGURIDAD EN UN DATA CENTER")
    
    add_h2("2.1 Representación Gráfica del Mapa Conceptual")
    add_p(
        "A continuación se presenta el mapa conceptual en alta resolución que organiza de manera jerárquica los componentes indispensables de la Seguridad Física y de la Seguridad Lógica, articulados bajo los principios de Confidencialidad, Integridad y Disponibilidad:"
    )

    img_hd_path = "mapa_conceptual_datacenter_hd.png"
    if os.path.exists(img_hd_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(6)
        r_img = p_img.add_run()
        r_img.add_picture(img_hd_path, width=Inches(6.6))
        
        p_fig = doc.add_paragraph()
        p_fig.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_fig.paragraph_format.space_after = Pt(8)
        r_fig = p_fig.add_run("Figura 1: Mapa Conceptual de Seguridad Física y Lógica en un Data Center (Estándares ISO 27001 / TIA-942)")
        r_fig.font.name = "Arial"
        r_fig.font.size = Pt(8.5)
        r_fig.font.italic = True
        r_fig.font.color.rgb = GRAY_TEXT

    add_h2("2.2 Desglose y Estructura Jerárquica del Mapa Conceptual")
    add_p(
        "Para una comprensión completa y exposición académica, el mapa conceptual se desglosa en los siguientes niveles jerárquicos:"
    )

    mapa_niveles = [
        ("Nivel 1 – Concepto Central:", "SEGURIDAD INTEGRAL EN UN DATA CENTER. Su objetivo es salvaguardar los activos de información garantizando la tríada CID: Confidencialidad (solo autorizados), Integridad (datos exactos y sin alteración) y Disponibilidad (servicios activos 24/7/365)."),
        ("Nivel 2 – Grandes Dimensiones:", "Se divide en dos áreas complementarias: 1) SEGURIDAD FÍSICA (protección tangible de instalaciones, equipos e infraestructura ambiental) y 2) SEGURIDAD LÓGICA (protección intangible de redes, sistemas operativos, bases de datos y accesos virtuales)."),
        ("Nivel 3A – Pilares de Seguridad Física:", "Comprende 5 componentes esenciales: a) Control de acceso perimetral; b) Videovigilancia y sensores de intrusión (CCTV); c) Climatización de precisión (HVAC); d) Supresión de incendios sin agua; y e) Respaldo y redundancia eléctrica (Tier III)."),
        ("Nivel 3B – Pilares de Seguridad Lógica:", "Comprende 5 componentes esenciales: a) Seguridad perimetral y segmentación de redes; b) Gestión de identidad y control de accesos (IAM); c) Monitoreo continuo y detección de intrusos (SOC/SIEM); d) Cifrado de datos en reposo y tránsito; y e) Continuidad de negocio y respaldos inmutables (DRP)."),
        ("Nivel 4 – Tecnologías y Controles Aplicados:", "Incluye herramientas específicas como Mantrap, Biometría, VESDA, gas Novec 1230, UPS modular, conmutador ATS, Firewalls NGFW, VLANs, DMZ, MFA/2FA, RBAC, IDS/IPS, SIEM, cifrado AES-256/TLS 1.3 y la regla de Backup 3-2-1 WORM.")
    ]

    for title, desc in mapa_niveles:
        add_bullet(title, desc)

    add_h2("2.3 Glosario Técnico de Términos del Mapa Conceptual")
    
    glosario_tbl = doc.add_table(rows=7, cols=2)
    glosario_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    glosario_tbl.autofit = False
    set_table_borders(glosario_tbl)

    glosario_data = [
        ("Esclusa (Mantrap)", "Sistema de control de acceso físico con dos puertas interbloqueadas donde una no puede abrirse si la otra no está totalmente cerrada."),
        ("VESDA", "Sistema de detección de humo por aspiración de muy alta sensibilidad que detecta partículas microscópicas antes de que aparezca llama abierta."),
        ("Novec 1230 / FM-200", "Gases limpios de extinción que sofocan el conato de fuego absorbiendo el calor sin dejar residuos, sin ser conductores de electricidad y sin dañar la electrónica."),
        ("Conmutador ATS", "Interruptor de transferencia automática que detecta la caída de la energía pública y enciende el generador diésel transfiriendo la carga sin intervención humana."),
        ("Firewall NGFW", "Cortafuegos de nueva generación que analiza el tráfico a nivel de aplicación (Capa 7), integrando antivirus, prevención de intrusos (IPS) y filtrado web."),
        ("Autenticación MFA / 2FA", "Mecanismo de seguridad que exige al menos dos factores distintos para validar la identidad (algo que sabe: contraseña, y algo que tiene: token/código)."),
        ("Inmutabilidad WORM", "Propiedad de almacenamiento (Write Once, Read Many) que impide que una copia de respaldo sea modificada, sobreescrita o borrada por un ataque de Ransomware.")
    ]

    for i, (termino, def_txt) in enumerate(glosario_data):
        row = glosario_tbl.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.0)
        c1.width = Inches(4.7)
        set_cell_background(c0, "F8FAFC")
        set_cell_background(c1, "FFFFFF")
        set_cell_padding(c0, top=50, bottom=50, left=80, right=80)
        set_cell_padding(c1, top=50, bottom=50, left=80, right=80)

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(termino)
        r0.font.name = "Arial"
        r0.font.size = Pt(9)
        r0.font.bold = True
        r0.font.color.rgb = NAVY

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(def_txt)
        r1.font.name = "Arial"
        r1.font.size = Pt(9)
        r1.font.color.rgb = DARK_TEXT

    # =========================================================================
    # CONCLUSIONES Y RECOMENDACIONES
    # =========================================================================
    add_h1("CONCLUSIONES Y RECOMENDACIONES GENERALES")
    
    conclusiones_list = [
        ("Indisociabilidad de Ambas Dimensiones",
         "La seguridad física y la seguridad lógica son dos caras de la misma moneda. Blindar la red con firewalls y cifrado resulta estéril si cualquier persona puede ingresar a la sala y desconectar un servidor; de igual modo, un centro acorazado con biometría es inútil si los puertos de administración están abiertos a Internet sin protección."),
        ("Criterio de Redundancia y Continuidad (Tier III)",
         "En centros de cómputo que manejan operaciones críticas, no debe existir ningún punto único de falla (SPOF - Single Point of Failure). La duplicidad de rutas eléctricas (UPS y generadores) y la climatización N+1 son requisitos indispensables para cumplir con los acuerdos de nivel de servicio (SLA)."),
        ("Evolución hacia la Inmutabilidad de los Datos",
         "Frente al crecimiento de amenazas avanzadas como el Ransomware corporativo, las copias de seguridad tradicionales son insuficientes. Es mandatorio implementar esquemas de respaldo inmutables (WORM) y almacenamiento fuera de línea bajo la regla 3-2-1."),
        ("Mantenimiento y Auditoría Periódica",
         "La seguridad de un Data Center no concluye con la compra de equipos; exige pruebas periódicas de arranque de generadores, simulacros de contingencia, revisión semanal de bitácoras de acceso y auditorías continuas bajo la norma ISO/IEC 27001.")
    ]

    for title, desc in conclusiones_list:
        add_bullet(title, desc)

    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_end.paragraph_format.space_before = Pt(20)
    r_end = p_end.add_run("— FIN DE LA ACTIVIDAD DE APRENDIZAJE —")
    r_end.font.name = "Arial"
    r_end.font.size = Pt(9.5)
    r_end.font.bold = True
    r_end.font.color.rgb = GRAY_TEXT

    output_main = "Caso_Practico_y_Mapa_Conceptual_DataCenter_Grupo.docx"
    doc.save(output_main)
    print(f"Documento guardado: {output_main}")

    output_alt = "TRABAJO_COMPLETO_DATA_CENTER_CASO_Y_MAPA.docx"
    doc.save(output_alt)
    print(f"Documento guardado: {output_alt}")

    try:
        doc.save("Caso_Practico_Seguridad_DataCenter_Grupo.docx")
        print("Documento guardado: Caso_Practico_Seguridad_DataCenter_Grupo.docx")
    except PermissionError:
        print("Nota: El archivo anterior estaba abierto en Word. Usa el nuevo archivo creado.")

if __name__ == "__main__":
    build_full_assignment()
